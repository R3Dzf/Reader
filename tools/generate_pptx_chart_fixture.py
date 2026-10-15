"""Independent python-pptx chart caches/relationships, not an application dependency."""
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.util import Inches,Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
p=Presentation();p.slide_width=Inches(10);p.slide_height=Inches(6)
def category(kind,title,categories,series):
 data=CategoryChartData();data.categories=categories
 for name,values in series:data.add_series(name,values)
 return add(kind,title,data)
def add(kind,title,data):
 slide=p.slides.add_slide(p.slide_layouts[6])
 slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
 chart=slide.shapes.add_chart(kind,Inches(.5),Inches(.3),Inches(9),Inches(5.4),data).chart
 chart.has_title=True;chart.chart_title.text_frame.text=title
 chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM
 for index,series in enumerate(chart.series):
  color=RGBColor.from_string(['1464C8','DA492A','469C35'][index%3])
  series.format.fill.solid();series.format.fill.fore_color.rgb=color
  series.format.line.color.rgb=color
 return chart
c=category(XL_CHART_TYPE.COLUMN_CLUSTERED,'مبيعات الأشهر',['يناير','فبراير','مارس'],[('المبيعات',[12,-6,18]),('التكلفة',[8,5,9])])
c.value_axis.minimum_scale=-10;c.value_axis.maximum_scale=20;c.value_axis.major_unit=5
c.plots[0].has_data_labels=True;c.plots[0].data_labels.show_value=True
c=category(XL_CHART_TYPE.BAR_STACKED,'مقارنة المناطق',['شرق','غرب','وسط'],[('أ',[10,20,30]),('ب',[5,-8,9])])
c=category(XL_CHART_TYPE.COLUMN_STACKED_100,'النسب المحفوظة',['شرق','غرب','وسط'],[('أ',[10,15,20]),('ب',[30,15,20])])
c.value_axis.minimum_scale=0;c.value_axis.maximum_scale=1;c.value_axis.tick_labels.number_format='0%'
c=category(XL_CHART_TYPE.LINE_MARKERS,'تغيّر القيم',['يناير','فبراير','مارس','أبريل'],[('أ',[5,None,12,9]),('ب',[4,8,None,11])])
c=category(XL_CHART_TYPE.AREA,'مساحة البيانات',['واحد','اثنان','ثلاثة'],[('التجربة',[3,9,5])])
c=category(XL_CHART_TYPE.PIE,'توزيع الأقسام',['قسم أ','قسم ب','قسم ج'],[('الحصة',[10,20,30])])
c.plots[0].has_data_labels=True;c.plots[0].data_labels.show_percentage=True;c.plots[0].data_labels.show_value=False
for point,color in zip(c.series[0].points,['1464C8','DA492A','469C35']):point.format.fill.solid();point.format.fill.fore_color.rgb=RGBColor.from_string(color)
c=category(XL_CHART_TYPE.DOUGHNUT,'حلقة القيم',['أ','ب','ج'],[('الدونات',[10,20,30])])
# python-pptx has no hole_size property; write the actual DrawingML value.
c._chartSpace.xpath('//c:holeSize')[0].set('val','60')
assert c._chartSpace.xpath('//c:holeSize')[0].get('val')=='60'
xy=XyChartData();s=xy.add_series('القياسات')
for x,y in [(1,2),(2,4),(10,3)]:s.add_data_point(x,y)
c=add(XL_CHART_TYPE.XY_SCATTER,'قياسات غير متساوية',xy)
c._chartSpace.xpath('//c:scatterStyle')[0].set('val','marker')
c.series[0].format.line.fill.background()
c.category_axis.minimum_scale=0;c.category_axis.maximum_scale=10
c=category(XL_CHART_TYPE.RADAR,'نوع غير مدعوم',['أ','ب','ج'],[('القيم',[2,4,3])])
# The writer creates real cached data and an embedded workbook. The reader uses caches only.
path=ROOT/'app/src/test/resources/fixtures/charts-native-arabic.pptx';p.save(path);print(path)
