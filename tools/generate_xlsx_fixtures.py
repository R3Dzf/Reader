"""Developer-only independent XLSX fixture generator (openpyxl 3.1.5).
Run: python tools/generate_xlsx_fixtures.py
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import Workbook

root = Path(__file__).resolve().parents[1]
target = root / "app/src/test/resources/fixtures/number-formats-arabic.xlsx"
target.parent.mkdir(parents=True, exist_ok=True)
book = Workbook()
sheet = book.active
sheet.title = "الميزانية"
sheet.sheet_view.rightToLeft = True
examples = [
    (1234.567, "#,##0.00"), (0.125, "0.0%"), (45292.5, "yyyy-mm-dd hh:mm:ss"),
    (-1234.5, '#,##0.00;[Red](#,##0.00)'), (1234.5, '"ج.م "#,##0.00'),
    (1.5, "[h]:mm:ss"), (60, "yyyy-mm-dd"), ("00123", "0.00"),
    (True, "0.00"), ("#DIV/0!", "0.00"), (0.333, "# ?/?"),
    ("=1/4", "0.00%"), ("=SUM(2,3)", "0.00"), (1234, "0.00E+00"),
]
for row, (value, code) in enumerate(examples, 1):
    sheet.cell(row, 1, value).number_format = code
    sheet.cell(row, 2, f"اختبار {row}")
book.save(target)
# A saved formula result, independently of Reader's parser; no formula calculation at runtime.
with ZipFile(target) as archive:
    parts = {name: archive.read(name) for name in archive.namelist()}
ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
xml = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
xml.find('.//s:c[@r="A12"]/s:v', ns).text = "0.25"
parts["xl/worksheets/sheet1.xml"] = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
with ZipFile(target, "w", ZIP_DEFLATED) as archive:
    for name, data in parts.items():
        archive.writestr(name, data)
print(target)

# A second workbook exercises real worksheet geometry and saved cell appearance.
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
book = Workbook()
sheet = book.active
sheet.title = "دمج وتخطيط"
sheet.sheet_view.rightToLeft = True
sheet.sheet_view.showGridLines = False
for column, width in [("A", 12), ("B", 24), ("C", 8), ("D", 16), ("E", 10), ("F", 20)]:
    sheet.column_dimensions[column].width = width
sheet.column_dimensions["E"].hidden = True
for row, height in [(1, 24), (2, 36), (3, 18), (4, 27), (6, 30)]:
    sheet.row_dimensions[row].height = height
sheet.row_dimensions[5].hidden = True
sheet.merge_cells("A1:B3")
sheet["A1"] = "عنوان عربي مدمج"
sheet["A1"].font = Font(name="Arial", size=16, bold=True, italic=True, underline="single", strike=True, color="112233")
sheet["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
sheet["A1"].border = Border(left=Side(style="thick", color="FF0000"), right=Side(style="double", color="0000FF"),
    top=Side(style="dashed", color="00AA00"), bottom=Side(style="dotted", color="222222"))
sheet.merge_cells("C1:C2")
sheet["C1"] = "عمودي"
sheet["C1"].alignment = Alignment(horizontal="right", vertical="top", readingOrder=2)
sheet.merge_cells("A4:D4")
sheet["A4"] = "أفقي"
sheet["D6"] = 1234.5
sheet["D6"].number_format = "#,##0.00"
sheet["F6"] = "Visible after hidden E"
sheet["B1000000"] = "صف بعيد"
book.save(root / "app/src/test/resources/fixtures/worksheet-layout-arabic.xlsx")
print(root / "app/src/test/resources/fixtures/worksheet-layout-arabic.xlsx")
