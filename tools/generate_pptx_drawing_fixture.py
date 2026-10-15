"""Independent python-pptx drawing/font/color fixture, with explicit DrawingML edge cases."""
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
p = Presentation(); p.slide_width=Inches(10); p.slide_height=Inches(5)
def box(slide,kind,x,y,w,h,color):
    s=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color);s.line.fill.background()
    return s
s=p.slides.add_slide(p.slide_layouts[6])
# Real master decoration: copied XML is independent of the Android model.
bar=box(s,MSO_SHAPE.RECTANGLE,0,0,10,.2,'124578')
p.slide_master._element.spTree.append(deepcopy(bar._element)); s.shapes._spTree.remove(bar._element)
layout=box(s,MSO_SHAPE.RECTANGLE,9,4,1,1,'228833')
p.slide_layouts[6]._element.spTree.append(deepcopy(layout._element)); s.shapes._spTree.remove(layout._element)
group=s.shapes.add_group_shape()
ellipse=group.shapes.add_shape(MSO_SHAPE.OVAL,0,0,Inches(1),Inches(1));ellipse.rotation=30
ellipse.fill.solid();ellipse.fill.fore_color.rgb=RGBColor(255,0,0)
ellipse.line.color.rgb=RGBColor(0,0,255);ellipse.line.width=Pt(2)
text=group.shapes.add_textbox(Inches(2),0,Inches(2),Inches(1))
text.text_frame.margin_left=Pt(0);text.text_frame.margin_right=Pt(0)
text.text_frame.margin_top=Pt(0);text.text_frame.margin_bottom=Pt(0)
from pptx.enum.text import MSO_ANCHOR
text.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
para=text.text_frame.paragraphs[0]
r=para.add_run();r.text='مرحبا';r.font.name='Cairo';r.font.size=Pt(24)
r=para.add_run();r.text=' Reader';r.font.name='Lato';r.font.size=Pt(20);r.font.bold=True
pr=para._p.get_or_add_pPr();pr.set('rtl','1');pr.set('marL',str(Pt(12)));pr.set('indent',str(Pt(-6)))
bu=OxmlElement('a:buChar');bu.set('char','•');pr.append(bu)
group.left=Inches(2);group.top=Inches(1);group.width=Inches(4);group.height=Inches(2)
# White style-matrix background; ancestor background is intentionally yellow.
masterbg=p.slide_master._element.get_or_add_cSld().get_or_add_bg() if False else None
for i in range(3):
    slide=p.slides.add_slide(p.slide_layouts[6]);slide._element.set('showMasterSp','0')
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
    if i==0:
        shape=box(slide,MSO_SHAPE.DIAMOND,1,1,2,2,'FFFFFF');shape.rotation=45
        shape._element.spPr.xfrm.set('flipH','1')
    if i==1:
        shape=box(slide,MSO_SHAPE.ROUNDED_RECTANGLE,1,1,8,3,'FFFFFF')
        fill=shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        shape._element.spPr.remove(fill)
        grad=OxmlElement('a:gradFill');lst=OxmlElement('a:gsLst');grad.append(lst)
        for position,color in [('0','0000FF'),('100000','FF0000')]:
            stop=OxmlElement('a:gs');stop.set('pos',position);rgb=OxmlElement('a:srgbClr');rgb.set('val',color)
            stop.append(rgb);lst.append(stop)
        lin=OxmlElement('a:lin');lin.set('ang','0');grad.append(lin);shape._element.spPr.append(grad)
path=ROOT/'app/src/test/resources/fixtures/drawing-fonts-white.pptx';p.save(path)
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with ZipFile(path) as z: entries={name:z.read(name) for name in z.namelist()}
root=etree.fromstring(entries['ppt/slides/slide1.xml']);csld=root.find('p:cSld',NS)
bg=etree.Element('{'+NS['p']+'}bg');ref=etree.SubElement(bg,'{'+NS['p']+'}bgRef',idx='1001');etree.SubElement(ref,'{'+NS['a']+'}schemeClr',val='bg1');csld.insert(0,bg)
entries['ppt/slides/slide1.xml']=etree.tostring(root)
# This inherited yellow MUST NOT leak through an explicit white bgRef on slide 1.
root=etree.fromstring(entries['ppt/slideMasters/slideMaster1.xml']);bg=root.find('p:cSld/p:bg',NS)
if bg is not None: bg.getparent().remove(bg)
bg=etree.Element('{'+NS['p']+'}bg');pr=etree.SubElement(bg,'{'+NS['p']+'}bgPr');fill=etree.SubElement(pr,'{'+NS['a']+'}solidFill');etree.SubElement(fill,'{'+NS['a']+'}srgbClr',val='FFFF00');root.find('p:cSld',NS).insert(0,bg)
entries['ppt/slideMasters/slideMaster1.xml']=etree.tostring(root)
with ZipFile(path,'w',ZIP_DEFLATED) as z:
 for name,data in entries.items():z.writestr(name,data)
print(path)
